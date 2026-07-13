import os
from bs4 import BeautifulSoup
from pypdf import PdfReader
from docx import Document as DocxReader
from pptx import Presentation as PptxReader
from odf import text,teletype
from odf.opendocument import load
import re

from vision import get_image_caption

def extract_text_factory(file_path: str) -> str:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Target file missing: {file_path}")
    
    _, ext = os.path.splitext(file_path)
    ext = ext.lower()
    document_name = os.path.basename(file_path)

    # 1. Plain Text & Markdown (.txt, .md)
    if ext in ['.txt', '.md']:
        with open(file_path,'r',encoding='utf-8') as f:
            return f.read()
    
    # 2. Modern Word Documents (.docx)
    elif ext == '.docx':
        doc = DocxReader(file_path)
        body_elements = []
        for para in doc.paragraphs:
            if para.text.strip():
                body_elements.append(para.text)

            if "w:drawing" in para._p.xml:
                # Extract image bytes via relationship ID if possible
                r_ids = re.findall(r'r:embed="([^"]+)"', para._p.xml)
                if r_ids:
                    for r_id in r_ids:
                        try:
                            image_part = doc.part.related_parts[r_id]
                            caption = get_image_caption(image_part.blob, document_name)
                            body_elements.append(caption)
                        except Exception:
                            body_elements.append("[IMAGE ALIAS: Embedded Word Diagram/Graphic Layout]")
                else:
                    body_elements.append("[IMAGE ALIAS: Embedded Word Diagram/Graphic Layout]")
        return "\n".join(body_elements)
    
    # 3. PowerPoint Presentations (.pptx)
    elif ext == '.pptx':
        prs = PptxReader(file_path)
        slide_text = []
        for i, slide in enumerate(prs.slides, start =1):
            for shape in slide.shapes:
                if hasattr(shape,"text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
                elif shape.shape_type == 13:
                    try:
                        caption = get_image_caption(shape.image.blob, document_name)
                        slide_text.append(f"\n{caption}\n")
                    except Exception:
                        slide_text.append(f"\n[IMAGE ALIAS: Presentation Visual Graphic on Slide {i}]\n")
        return "\n".join(slide_text)
    
    # 4. Portable Document Format (.pdf)
    elif ext == '.pdf':
        reader = PdfReader(file_path)
        extracted_pages = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                extracted_pages.append(page_text)
            
            if page.images:
                for img in page.images:
                    try:
                        caption = get_image_caption(img.data, document_name)
                        extracted_pages.append(f"\n{caption}\n")
                    except Exception:
                        extracted_pages.append("\n[IMAGE ALIAS: Embedded PDF Visual Asset/Scan]\n")
        return "\n".join(extracted_pages)
    
    # 5. HTML / Wiki Portals (.html, .htm)
    elif ext in ['.html', '.htm']:
        with open(file_path, 'r', encoding='utf-8') as f:
            soup = BeautifulSoup(f.read(), 'html.parser')
            # Extract and swap images with their semantic alt-text
            for img in soup.find_all('img'):
                alt_text = img.get('alt', '').strip()
                semantic_token = f" [IMAGE ALT-TEXT: {alt_text}] " if alt_text else " [IMAGE ALIAS: Unlabeled Web Graphic] "
                img.replace_with(semantic_token)
            
            # Strip formatting and return pure text
            for script in soup(["script", "style"]):
                script.decompose()
            return "\n".join([line.strip() for line in soup.get_text().splitlines() if line.strip()])

    # 6. OpenDocument Text (.odt)
    elif ext == '.odt':
        odt_doc = load(file_path)
        paragraphs = odt_doc.getElementsByType(text.P)
        return "\n".join([teletype.extractText(p) for p in paragraphs if teletype.extractText(p).strip()])
    
    else:
        raise ValueError(f"Unsupported file format:{ext}")
